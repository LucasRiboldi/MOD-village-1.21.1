"""Exporta os blocos dos templates de vila Vanilla e suas rotas no mod.

Uso:
  python scripts/export_vanilla_village_block_routes.py <minecraft-server.jar> <saida.xlsx>
"""

from __future__ import annotations

import gzip
import re
import struct
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


TAG_END = 0
TAG_BYTE = 1
TAG_SHORT = 2
TAG_INT = 3
TAG_LONG = 4
TAG_FLOAT = 5
TAG_DOUBLE = 6
TAG_BYTE_ARRAY = 7
TAG_STRING = 8
TAG_LIST = 9
TAG_COMPOUND = 10
TAG_INT_ARRAY = 11
TAG_LONG_ARRAY = 12

IGNORED_TEMPLATE_BLOCKS = {
    "minecraft:air",
    "minecraft:cave_air",
    "minecraft:void_air",
    "minecraft:jigsaw",
    "minecraft:structure_block",
    "minecraft:structure_void",
}

PROFESSION_FOR_PRODUCTION = {
    "HARVESTED": "LENHADOR",
    "FARMED": "FAZENDEIRO",
    "MINED": "MINEIRO",
    "SURFACE_GATHERED": "FUNDIDOR",
    "SOIL_GATHERED": "FAZENDEIRO",
    "SHEARED": "PASTOR",
    "CRAFTED_WOOD": "CARPINTEIRO",
    "CRAFTED_STONE": "PEDREIRO",
    "SMELTED": "FUNDIDOR",
}

GROUND_SHAPING = {
    "farmland",
    "water",
    "dirt_path",
    "wheat",
    "carrots",
    "potatoes",
    "beetroots",
}

WOOD_SUFFIXES = (
    "_planks",
    "_stairs",
    "_slab",
    "_fence",
    "_fence_gate",
    "_door",
    "_trapdoor",
    "_pressure_plate",
    "_button",
    "_sign",
    "_hanging_sign",
)


class NbtReader:
    """Leitor pequeno para o formato NBT usado pelos templates de estrutura."""

    def __init__(self, data: bytes):
        self.data = memoryview(data)
        self.position = 0

    def read(self, fmt: str):
        size = struct.calcsize(fmt)
        value = struct.unpack_from(fmt, self.data, self.position)
        self.position += size
        return value[0] if len(value) == 1 else value

    def string(self) -> str:
        length = self.read(">H")
        value = bytes(self.data[self.position:self.position + length]).decode("utf-8")
        self.position += length
        return value

    def payload(self, tag: int):
        if tag == TAG_END:
            return None
        if tag == TAG_BYTE:
            return self.read(">b")
        if tag == TAG_SHORT:
            return self.read(">h")
        if tag == TAG_INT:
            return self.read(">i")
        if tag == TAG_LONG:
            return self.read(">q")
        if tag == TAG_FLOAT:
            return self.read(">f")
        if tag == TAG_DOUBLE:
            return self.read(">d")
        if tag == TAG_BYTE_ARRAY:
            length = self.read(">i")
            value = bytes(self.data[self.position:self.position + length])
            self.position += length
            return value
        if tag == TAG_STRING:
            return self.string()
        if tag == TAG_LIST:
            child_type = self.read(">b")
            length = self.read(">i")
            return [self.payload(child_type) for _ in range(length)]
        if tag == TAG_COMPOUND:
            result = {}
            while True:
                child_type = self.read(">b")
                if child_type == TAG_END:
                    return result
                key = self.string()
                result[key] = self.payload(child_type)
        if tag == TAG_INT_ARRAY:
            length = self.read(">i")
            return [self.read(">i") for _ in range(length)]
        if tag == TAG_LONG_ARRAY:
            length = self.read(">i")
            return [self.read(">q") for _ in range(length)]
        raise ValueError(f"Unsupported NBT tag {tag}")

    def root(self):
        tag = self.read(">b")
        if tag != TAG_COMPOUND:
            raise ValueError(f"Template root must be a compound, got {tag}")
        self.string()
        return self.payload(tag)


def direct_resource_routes(resource_type_file: Path) -> dict[str, str]:
    """Lê Production de ResourceType para não duplicar o catálogo Java."""
    source = resource_type_file.read_text(encoding="utf-8")
    enum_body = source.split("public enum ResourceType {", 1)[1].split("private final", 1)[0]
    routes = {}

    for chunk in enum_body.split("),"):
        name = re.search(r"^\s*([A-Z][A-Z0-9_]*)\(", chunk, re.MULTILINE)
        production = re.search(r"Production\.([A-Z_]+)", chunk)
        if name and production:
            routes[name.group(1).lower()] = PROFESSION_FOR_PRODUCTION[production.group(1)]

    return routes


def route_for(block_id: str, direct_routes: dict[str, str]) -> str:
    path = block_id.removeprefix("minecraft:")

    if path in direct_routes:
        return direct_routes[path]
    if path in GROUND_SHAPING:
        return "CONSTRUTOR - molda no local"
    if path == "lava":
        return "SEM ROTA - lava nao e colocada pela colonia"
    if path.startswith("potted_"):
        return "CONSTRUTOR - monta vaso e planta no local"
    if path.endswith(WOOD_SUFFIXES):
        return "CARPINTEIRO - receita de madeira do LENHADOR"
    if path.endswith("_log") or path.endswith("_wood"):
        return "LENHADOR - coleta madeira"
    if path.endswith(("_leaves", "_sapling", "_bush", "_flower")):
        return "SEM ROTA DIRETA - recurso natural"
    if any(token in path for token in ("stone", "sandstone", "cobblestone", "brick", "terracotta", "glass")):
        return "PEDREIRO - receita com insumo do MINEIRO/FUNDIDOR"
    return "SEM ROTA DIRETA - sem profissao que adquira"


def template_blocks(template_data: bytes) -> set[str]:
    if template_data.startswith(b"\x1f\x8b"):
        template_data = gzip.decompress(template_data)
    root = NbtReader(template_data).root()
    palettes = root.get("palettes") or [root.get("palette", [])]
    states = {entry.get("state") for entry in root.get("blocks", [])}
    blocks = set()

    for palette in palettes:
        for index, state in enumerate(palette):
            if index not in states:
                continue
            name = state.get("Name")
            if name and name not in IGNORED_TEMPLATE_BLOCKS:
                blocks.add(name)

    return blocks


def village_blocks(minecraft_jar: Path) -> set[str]:
    prefix = "data/minecraft/structure/village/"
    with zipfile.ZipFile(minecraft_jar) as jar:
        names = [name for name in jar.namelist() if name.startswith(prefix) and name.endswith(".nbt")]
        if not names:
            raise ValueError(f"No village templates found in {minecraft_jar}")

        blocks = set()
        for name in names:
            try:
                blocks.update(template_blocks(jar.read(name)))
            except (OSError, ValueError, struct.error) as error:
                raise ValueError(f"Could not read village template {name}") from error

        return blocks


def write_sheet(output: Path, blocks: set[str], direct_routes: dict[str, str]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blocos de vila Vanilla"
    sheet.append(["Bloco", "Profissao/Rota"])

    for block_id in sorted(blocks):
        sheet.append([block_id, route_for(block_id, direct_routes)])

    header_fill = PatternFill("solid", fgColor="4F6228")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 58
    workbook.save(output)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)

    minecraft_jar = Path(sys.argv[1])
    output = Path(sys.argv[2])
    resource_type_file = Path("src/main/java/com/villagecolony/core/type/ResourceType.java")
    blocks = village_blocks(minecraft_jar)
    write_sheet(output, blocks, direct_resource_routes(resource_type_file))
    print(f"{len(blocks)} blocks exported to {output}")


if __name__ == "__main__":
    main()
